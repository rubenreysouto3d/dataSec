#!/usr/bin/env python3
"""Validate the official London ward population source used for resident context."""

from __future__ import annotations

import io
import json
import re
import urllib.request

from openpyxl import load_workbook

PACKAGE_API = (
    "https://data.london.gov.uk/api/action/package_show"
    "?id=2021-census-wards-demography-and-migration-vqlx7"
)
RESOURCE_NAME = "Usual Residents.xlsx"
WARD_CODE = re.compile(r"^E050\d{5}$")
KNOWN_DATASEC_CODES = {"E05009317", "E05009318"}


def fetch(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "dataSec-source-health/0.3"},
    )
    with urllib.request.urlopen(request, timeout=45) as response:
        return response.read()


package = json.loads(fetch(PACKAGE_API))
if not package.get("success"):
    raise RuntimeError("London population CKAN package lookup failed")

resources = package.get("result", {}).get("resources", [])
resource = next(
    (
        item
        for item in resources
        if str(item.get("name") or item.get("description") or "").strip() == RESOURCE_NAME
    ),
    None,
)
if resource is None:
    raise RuntimeError(f"London population resource not found: {RESOURCE_NAME}")

url = str(resource.get("url") or "").strip()
if url.startswith("http://"):
    url = "https://" + url.removeprefix("http://")
if not url.startswith("https://"):
    raise RuntimeError(f"London population resource has no usable URL: {url!r}")

payload = fetch(url)
if len(payload) < 50_000:
    raise RuntimeError(f"Implausibly small London population workbook: {len(payload):,} bytes")

workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
codes: set[str] = set()
for worksheet in workbook.worksheets:
    for row in worksheet.iter_rows(values_only=True):
        for value in row:
            text = str(value or "").strip()
            if WARD_CODE.fullmatch(text):
                codes.add(text)

missing_known = sorted(KNOWN_DATASEC_CODES - codes)
if missing_known:
    raise RuntimeError(
        "Official population workbook does not contain known dataSec ward codes: "
        + ", ".join(missing_known)
    )
if len(codes) < 600:
    raise RuntimeError(f"Implausibly low London ward-code coverage: {len(codes)}")

print(
    json.dumps(
        {
            "ok": True,
            "resource": RESOURCE_NAME,
            "resourceId": resource.get("id"),
            "workbookBytes": len(payload),
            "worksheets": workbook.sheetnames,
            "wardCodeCount": len(codes),
            "knownDataSecCodesPresent": sorted(KNOWN_DATASEC_CODES),
        },
        indent=2,
    )
)
