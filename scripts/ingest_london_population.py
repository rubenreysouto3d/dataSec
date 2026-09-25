#!/usr/bin/env python3
"""Ingest 2021 Census usual residents for London wards."""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import time
import urllib.parse
import urllib.request
import uuid

from openpyxl import load_workbook

try:
    from scripts.ingest_london import SupabaseRest, USER_AGENT
except ModuleNotFoundError:
    from ingest_london import SupabaseRest, USER_AGENT

PACKAGE_API = "https://data.london.gov.uk/api/action/package_show?id=vqlx7"
DATASET_URL = "https://data.london.gov.uk/dataset/2021-census-wards-demography-and-migration-vqlx7"
RESOURCE_NAME = "Usual Residents.xlsx"
SOURCE_SLUG = "london-census-2021-ward-population"
CENSUS_DATE = "2021-03-21"
WARD_CODE = re.compile(r"^E050\d{5}$")


def log(message: str) -> None:
    print(message, flush=True)


def fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def population_resource() -> tuple[str, str]:
    package = json.loads(fetch(PACKAGE_API))
    if not package.get("success"):
        raise RuntimeError("London Datastore package lookup failed")
    resources = package.get("result", {}).get("resources", [])
    resource = next(
        (
            item
            for item in resources
            if str(item.get("name") or item.get("description") or "").strip() == RESOURCE_NAME
        ),
        None,
    )
    if resource is None:
        raise RuntimeError(f"London population resource not found: {RESOURCE_NAME}")

    url = str(resource.get("url") or "").strip()
    if url.startswith("http://"):
        url = "https://" + url.removeprefix("http://")
    if not url.startswith("https://"):
        raise RuntimeError(f"London population resource has no usable URL: {url!r}")
    return url, str(resource.get("id") or "")


def parse_population(blob: bytes) -> dict[str, int]:
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if "2021" not in workbook.sheetnames:
        raise RuntimeError("London population workbook no longer contains a 2021 sheet")

    sheet = workbook["2021"]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(value or "").strip() for value in next(rows)]
    except StopIteration as exc:
        raise RuntimeError("London population sheet is empty") from exc

    required = {"ward code", "All usual residents"}
    missing = required - set(header)
    if missing:
        raise RuntimeError(f"London population schema changed; missing: {sorted(missing)}")

    code_index = header.index("ward code")
    population_index = header.index("All usual residents")
    totals: dict[str, int] = {}

    for row in rows:
        code = str(row[code_index] or "").strip()
        if not WARD_CODE.fullmatch(code):
            continue
        raw_population = row[population_index]
        if isinstance(raw_population, bool) or not isinstance(raw_population, (int, float)):
            raise RuntimeError(f"Invalid population value for {code}: {raw_population!r}")
        population = int(raw_population)
        if population <= 0:
            raise RuntimeError(f"Non-positive population for {code}: {population}")
        if code in totals:
            raise RuntimeError(f"Duplicate ward code in population workbook: {code}")
        totals[code] = population

    if len(totals) != 679:
        raise RuntimeError(f"Expected 679 London ward populations, got {len(totals)}")

    city_total = sum(totals.values())
    if not 8_000_000 <= city_total <= 10_000_000:
        raise RuntimeError(f"Implausible London 2021 population total: {city_total:,}")
    return totals



def persist(db: SupabaseRest, resource_id: str, totals: dict[str, int]) -> None:
    # Source-health independently validates that the official workbook exposes
    # the same 679 E050 ward identifiers used by dataSec. Database foreign keys
    # fail closed if any source code is not a stored area.
    db.upsert(
        "sources",
        [{
            "slug": SOURCE_SLUG,
            "authority": "Greater London Authority / Office for National Statistics",
            "source_url": DATASET_URL,
            "licence": "Open Government Licence v3.0",
            "update_frequency": "decennial",
            "source_type": "census_population",
            "granularity": "ward",
            "notes": (
                "2021 Census usual residents by ward. Used as a resident denominator/context; "
                "the population vintage is older than current crime observations."
            ),
        }],
        "slug",
    )

    run_id = str(uuid.uuid4())
    db.upsert(
        "ingestion_runs",
        [{
            "id": run_id,
            "source_slug": SOURCE_SLUG,
            "status": "running",
            "source_version": "2021 Census",
            "row_count": len(totals),
            "matched_row_count": len(totals),
            "diagnostics": {
                "resource_id": resource_id,
                "ward_count": len(totals),
                "usual_residents": sum(totals.values()),
                "census_date": CENSUS_DATE,
                "measurement": "usual_residents",
            },
        }],
        "id",
    )

    try:
        rows = [
            {
                "area_id": f"gb-london-metropolitan:{code}",
                "source_slug": SOURCE_SLUG,
                "period_start": CENSUS_DATE,
                "population": population,
                "provenance": {
                    "resource_id": resource_id,
                    "worksheet": "2021",
                    "measurement": "usual_residents",
                    "census_date": CENSUS_DATE,
                    "source_vintage": "2021 Census",
                },
            }
            for code, population in sorted(totals.items())
        ]
        db.upsert(
            "area_population_snapshots",
            rows,
            "area_id,source_slug,period_start",
            batch=300,
        )

        query = urllib.parse.urlencode({"id": f"eq.{run_id}"})
        db.request(
            "ingestion_runs",
            method="PATCH",
            query=query,
            payload={
                "status": "passed",
                "finished_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            },
            prefer="return=minimal",
        )
    except Exception as exc:
        query = urllib.parse.urlencode({"id": f"eq.{run_id}"})
        try:
            db.request(
                "ingestion_runs",
                method="PATCH",
                query=query,
                payload={
                    "status": "failed",
                    "finished_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
                    "diagnostics": {"error": str(exc), "resource_id": resource_id},
                },
                prefer="return=minimal",
            )
        finally:
            raise


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    source_url, resource_id = population_resource()
    blob = fetch(source_url)
    totals = parse_population(blob)
    log(
        f"Validated London 2021 Census population: {len(totals)} wards; "
        f"{sum(totals.values()):,} usual residents"
    )

    if args.dry_run:
        log("Dry run complete; no database writes.")
        return 0

    supabase_url = os.environ.get("SUPABASE_URL")
    backend_key = os.environ.get("SUPABASE_SECRET_KEY") or os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url:
        raise RuntimeError("SUPABASE_URL is required")
    db = SupabaseRest(supabase_url, backend_key)
    persist(db, resource_id, totals)
    log("London 2021 Census population snapshot persisted successfully")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
